from pathlib import Path

from docx import Document as DocxDocument

from semantic_search_demo.models import Document


def load_docx(path: Path) -> Document:
    docx = DocxDocument(path)

    paragraphs: list[str] = []

    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    document_text = "\n\n".join(paragraphs)

    return Document(
        document_id=path.stem,
        source=str(path),
        file_type="docx",
        text=document_text,
    )